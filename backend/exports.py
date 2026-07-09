import io
import json
import csv
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ── CSV ──────────────────────────────────────────────────────────────────

def export_price_history_csv(products, history_rows) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Product", "Site", "URL", "Price (€)", "Was Price (€)", "Discount %", "Availability", "Scraped At"])
    product_map = {p.id: p for p in products}   # O(1) lookups; the linear scan per row was quadratic
    for h in history_rows:
        p = product_map.get(h.product_id)
        disc = ""
        if h.price and h.old_price and h.old_price > h.price + 0.5:
            disc = round((1 - h.price / h.old_price) * 100, 1)
        writer.writerow([
            p.name if p else "",
            p.site if p else "",
            p.url if p else "",
            h.price or "",
            h.old_price or "",
            disc,
            h.availability or "",
            h.scraped_at.strftime("%Y-%m-%d %H:%M") if h.scraped_at else "",
        ])
    return output.getvalue().encode("utf-8-sig")  # utf-8-sig for Excel compat


def export_alerts_csv(alerts) -> bytes:
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Product", "Site", "Condition", "Threshold (€)", "Price at Alert (€)", "Triggered At", "Email Sent"])
    for a in alerts:
        writer.writerow([
            a.product_name, a.site, a.condition,
            a.threshold, a.current_price,
            a.triggered_at.strftime("%Y-%m-%d %H:%M") if a.triggered_at else "",
            "Yes" if a.email_sent else "No",
        ])
    return output.getvalue().encode("utf-8-sig")


def export_competitors_csv(products, latest_prices: dict) -> bytes:
    """latest_prices: {product_id: {site: price}}"""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Product", "Kotsovolos (€)", "Public.gr (€)", "Plaisio (€)", "Best Price (€)", "Best Site"])
    for p in products:
        prices = latest_prices.get(p.id, {})
        vals = {k: v for k, v in prices.items() if v is not None}
        best_site = min(vals, key=vals.get) if vals else ""
        best_price = vals[best_site] if best_site else ""
        writer.writerow([
            p.name,
            prices.get("kotsovolos", ""),
            prices.get("public", ""),
            prices.get("plaisio", ""),
            best_price,
            best_site,
        ])
    return output.getvalue().encode("utf-8-sig")


# ── EXCEL ─────────────────────────────────────────────────────────────────

def _style_header_row(ws, row_num: int, num_cols: int):
    fill = PatternFill("solid", fgColor="0D0D0D")
    font = Font(color="C8F060", bold=True, name="Courier New", size=10)
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="left", vertical="center")


def _auto_width(ws):
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)


def export_full_excel(products, history_rows, alert_rows) -> bytes:
    wb = openpyxl.Workbook()

    # ── Sheet 1: Price History ──
    ws1 = wb.active
    ws1.title = "Price History"
    headers1 = ["Product", "Site", "Price (€)", "Was Price (€)", "Discount %", "Availability", "Scraped At"]
    ws1.append(headers1)
    _style_header_row(ws1, 1, len(headers1))
    ws1.row_dimensions[1].height = 20

    product_map = {p.id: p for p in products}
    for h in history_rows:
        p = product_map.get(h.product_id)
        disc = ""
        if h.price and h.old_price and h.old_price > h.price + 0.5:
            disc = round((1 - h.price / h.old_price) * 100, 1)
        ws1.append([
            p.name if p else "",
            p.site if p else "",
            h.price,
            h.old_price,
            disc,
            h.availability or "",
            h.scraped_at.strftime("%Y-%m-%d %H:%M") if h.scraped_at else "",
        ])
    _auto_width(ws1)

    # ── Sheet 2: Competitor Comparison ──
    ws2 = wb.create_sheet("Competitor Comparison")
    headers2 = ["Product", "Kotsovolos (€)", "Public.gr (€)", "Plaisio (€)", "Cheapest Site", "Best Price (€)", "Your Saving"]
    ws2.append(headers2)
    _style_header_row(ws2, 1, len(headers2))

    # Group latest price per product per site
    latest: dict[int, dict[str, float]] = {}
    for h in sorted(history_rows, key=lambda x: x.scraped_at or datetime.min):
        p = product_map.get(h.product_id)
        if p and h.price:
            latest.setdefault(p.id, {})[p.site] = h.price

    for p in products:
        prices = latest.get(p.id, {})
        vals = {k: v for k, v in prices.items() if v}
        best_site = min(vals, key=vals.get) if vals else ""
        best_price = vals.get(best_site, "")
        your_price = vals.get(p.site, "")
        saving = round(your_price - best_price, 2) if your_price and best_price else ""
        ws2.append([
            p.name,
            prices.get("kotsovolos", ""),
            prices.get("public", ""),
            prices.get("plaisio", ""),
            best_site,
            best_price,
            saving,
        ])
    _auto_width(ws2)

    # ── Sheet 3: Alert Log ──
    ws3 = wb.create_sheet("Alert Log")
    headers3 = ["Product", "Site", "Condition", "Threshold (€)", "Price at Alert (€)", "Triggered At", "Email Sent"]
    ws3.append(headers3)
    _style_header_row(ws3, 1, len(headers3))
    for a in alert_rows:
        ws3.append([
            a.product_name, a.site, a.condition,
            a.threshold, a.current_price,
            a.triggered_at.strftime("%Y-%m-%d %H:%M") if a.triggered_at else "",
            "Yes" if a.email_sent else "No",
        ])
    _auto_width(ws3)

    # ── Sheet 4: Products ──
    ws4 = wb.create_sheet("Products")
    headers4 = ["ID", "Name", "Site", "URL", "Alert Threshold (€)", "Alert Email", "Active", "Added"]
    ws4.append(headers4)
    _style_header_row(ws4, 1, len(headers4))
    for p in products:
        ws4.append([
            p.id, p.name, p.site, p.url,
            p.alert_threshold, p.alert_email,
            "Yes" if p.active else "No",
            p.created_at.strftime("%Y-%m-%d") if p.created_at else "",
        ])
    _auto_width(ws4)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
